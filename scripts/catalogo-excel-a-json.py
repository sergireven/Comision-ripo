"""Convierte la plantilla de artículos (Excel) a JSON para `npm run import-catalog`.

Uso:  python scripts/catalogo-excel-a-json.py plantilla-articulos-v2.xlsx catalogo.json
Necesita openpyxl (pip install openpyxl).
"""
import json
import sys

from openpyxl import load_workbook

ART_COLS = {
    "ID": "id", "Nombre (castellano)": "name", "Nombre (catalán)": "name_ca", "Precio (€)": "price",
    "Categoría": "category", "Colores (separados por comas)": "colors", "Opciones (separadas por comas)": "options",
    "Tallas (separadas por comas)": "sizes", "Personalización (qué debe indicar la familia)": "personalization",
    "Incluido en packs (ID de pack, separados por comas)": "packs", "Descripción (castellano)": "description",
    "Descripción (catalán)": "description_ca", "Foto (carpeta img)": "image",
    "Guía de tallas (carpeta img)": "size_guide", "Visible (Sí/No)": "visible",
}
PACK_COLS = {
    "ID": "id", "Nombre (castellano)": "name", "Nombre (catalán)": "name_ca", "Precio pack (€)": "price",
    "Regalo (ID de artículo, separados por comas)": "gifts", "Descripción (castellano)": "description",
    "Descripción (catalán)": "description_ca", "Foto (carpeta img)": "image", "Visible (Sí/No)": "visible",
}


def read(ws, cols):
    header = [c.value for c in ws[1]]
    missing = [h for h in cols if h not in header]
    if missing:
        sys.exit(f"Faltan columnas en la hoja «{ws.title}»: {', '.join(missing)}")
    rows = []
    for values in ws.iter_rows(min_row=2, values_only=True):
        row = {cols[h]: v for h, v in zip(header, values) if h in cols}
        if row.get("id") in (None, "") or not row.get("name"):
            continue
        rows.append({k: (v.strip() if isinstance(v, str) else v) for k, v in row.items()})
    return rows


def main():
    src, out = sys.argv[1], sys.argv[2]
    wb = load_workbook(src, data_only=True)
    data = {"articles": read(wb["Artículos"], ART_COLS), "packs": read(wb["Packs"], PACK_COLS)}
    with open(out, "w", encoding="utf8") as f:
        json.dump(data, f, ensure_ascii=False, indent=1)
    print(f"{len(data['articles'])} artículos y {len(data['packs'])} packs -> {out}")


if __name__ == "__main__":
    main()
